"""Portable Excel export for the standalone collector (same stack as DMP)."""
import json
import os
from decimal import Decimal, InvalidOperation
from pathlib import Path

ROOT_COLUMNS = [('rootWord','词根'),('searchPeople','搜索人数'),('searchPeopleGrow','搜索人数环比'),
                ('searchPeopleGrowYoy','搜索人数同比'),('continuousGrowPeriod','连续增长周期'),
                ('demand','需求分类编码'),('newRoot','是否新词根')]
KEY_COLUMNS = [('keyWord','关键词'),('rank','排名'),('searchPeople','搜索人数'),('clickPeople','点击人数'),
               ('clickRate','点击率'),('dealPeople','支付人数'),('dealRate','支付转化率'),
               ('searchPeopleRange','搜索人数范围'),('dealRateRange','支付转化率范围'),
               ('throughTrainPrice','直通车价格'),('searchGrow','searchGrow（原字段）'),
               ('doubleHighWord','doubleHighWord（原字段）'),('demand','需求分类编码'),
               ('gxb','gxb（原字段）'),('unknownMap','扩展字段（原始JSON）')]
PERCENT = {'searchPeopleGrow','searchPeopleGrowYoy','clickRate','dealRate'}
NUMBER = PERCENT | {'searchPeople','clickPeople','dealPeople','rank','continuousGrowPeriod','throughTrainPrice'}


def cell_value(key, value):
    if isinstance(value, (dict,list)):
        return json.dumps(value,ensure_ascii=False,separators=(',',':'))
    if isinstance(value,str) and key in NUMBER and value.strip():
        try:
            number = Decimal(value)
            if number.is_finite():
                # Excel stores only 15 significant digits; longer values stay text.
                if len(number.as_tuple().digits) > 15:
                    return value
                return int(number) if number == number.to_integral() else float(number)
        except InvalidOperation:
            pass
    return value


def export_result(result, output):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    if result.get('status') != 'complete' or not isinstance(result.get('keywords'),dict):
        raise ValueError('关键词尚未采集完成，不能导出完整 Excel。')
    output = Path(output)
    output.parent.mkdir(parents=True,exist_ok=True)
    temp = output.with_name(output.stem + '.tmp.xlsx')
    wb = Workbook()
    wb.remove(wb.active)
    jobs = [('关键词分析',result['keywords']['hotWordList'],KEY_COLUMNS),
            ('词根归类',result['roots']['keyWordVoList'],ROOT_COLUMNS)]
    for sheet_name,rows,known_cols in jobs:
        keys = {key for key,_ in known_cols}
        extras = sorted({k for r in rows for k in r} - keys)
        cols = known_cols + [(k,k+'（新增原字段）') for k in extras]
        sheet = wb.create_sheet(sheet_name)
        sheet.sheet_view.showGridLines = False
        sheet.cell(2,1,sheet_name).font = Font(name='微软雅黑',size=14,bold=True,color='203B61')
        period = result['period']
        sheet.cell(3,1,f"{result['selection']['category_name']}。{period['start']} 至 {period['end']}。共 {len(rows)} 条。")
        sheet.cell(4,1,f"来源：{result['source']}。抓取时间：{result['fetched_at']}。")
        limit = len(result['keywords']['hotWordList']) >= 2000
        sheet.cell(5,1,'保留接口返回的原始数值。关键词达到2000条，可能存在服务端上限。' if limit else '保留接口本次返回的全部记录和原始数值。')
        for j,(_,label) in enumerate(cols,1):
            c=sheet.cell(7,j,label)
            c.font=Font(name='微软雅黑',size=10,bold=True,color='FFFFFF')
            c.fill=PatternFill('solid',fgColor='203B61')
            c.alignment=Alignment(horizontal='center',vertical='center')
            sheet.column_dimensions[get_column_letter(j)].width = 34 if j==1 else max(18,min(36,len(label)*2+2))
        sheet.row_dimensions[7].height=25
        for i,row in enumerate(rows,8):
            for j,(key,_) in enumerate(cols,1):
                value=cell_value(key,row.get(key))
                if isinstance(value,str) and len(value)>32767:
                    raise ValueError('存在超过 Excel 单元格长度限制的字段，完整数据已保留在 JSON。')
                c=sheet.cell(i,j,value)
                if isinstance(value,str):
                    c.data_type='s'  # literal keyword text, including strings beginning '='
                c.font=Font(name='微软雅黑',size=10)
                c.alignment=Alignment(vertical='center',horizontal='right' if isinstance(value,(int,float)) and not isinstance(value,bool) else 'left')
                if isinstance(value,(int,float)) and not isinstance(value,bool):
                    c.number_format='0.00%' if key in PERCENT else '#,##0.############'
        if rows:
            t=Table(displayName='KeywordData' if sheet_name=='关键词分析' else 'RootData',ref=f'A7:{get_column_letter(len(cols))}{len(rows)+7}')
            t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            sheet.add_table(t)
        sheet.freeze_panes='B8'
    try:
        wb.save(temp)
        wb.close()
        verify=load_workbook(temp,read_only=True,data_only=True)
        try:
            for name,rows,_ in jobs:
                if verify[name].max_row != len(rows)+7:
                    raise ValueError('导出的 Excel 行数校验失败。')
        finally:
            verify.close()
        os.replace(temp,output)
    except Exception:
        if temp.exists():
            temp.unlink()
        raise
    return output
