"""Batch planning, recoverable progress and spreadsheet export."""

import calendar
import json
import os
import re
import tempfile
import time
import zipfile
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from coverage import ROOT, build_payload, date_value


def periods_between(start, end, monthly=True):
    first = date.fromisoformat(date_value(start)[:4] + "-" + date_value(start)[4:6] + "-" + date_value(start)[6:])
    last = date.fromisoformat(date_value(end)[:4] + "-" + date_value(end)[4:6] + "-" + date_value(end)[6:])
    if first > last:
        raise ValueError("开始日期不能晚于结束日期")
    if not monthly:
        return [{"start_date": first.isoformat(), "end_date": last.isoformat()}]
    result = []
    current = first
    while current <= last:
        month_end = date(current.year, current.month, calendar.monthrange(current.year, current.month)[1])
        result.append({"start_date": current.isoformat(), "end_date": min(month_end, last).isoformat()})
        current = date(current.year + (current.month == 12), current.month % 12 + 1, 1)
    return result


def make_plan(settings):
    categories = settings["categories"]
    bands = settings["price_ranges"]
    category_bands = settings.get("category_price_ranges", {})
    periods = settings["periods"]
    mode = settings.get("query_mode")
    if mode is not None and mode not in {"first", "second", "third", "mixed"}:
        raise ValueError("查询方式必须为一级汇总、二级明细、三级明细或二三级混合")
    expected_level = {"first": 1, "second": 2, "third": 3}.get(mode, 2)
    if mode == "mixed" and {c.get("level") for c in categories} != {2, 3}:
        raise ValueError("混合查询须同时包含二级和三级类目")
    if mode != "mixed" and mode is not None and any(c.get("level", expected_level) != expected_level for c in categories):
        raise ValueError("所选类目层级与查询方式不一致，请重新选择")
    if not categories or not bands or not periods:
        raise ValueError("请至少填写一个类目、一个金额区间和一个日期区间")
    if not isinstance(category_bands, dict) or set(category_bands) - {str(c["category_id"]) for c in categories}:
        raise ValueError("类目金额模板与所选类目不匹配")
    if len({str(item["category_id"]) for item in categories}) != len(categories):
        raise ValueError("类目 ID 不能重复")
    if any(not str(item["category_name"]).strip() for item in categories):
        raise ValueError("类目名称不能为空")
    for category in categories:
        selected_bands = category_bands.get(str(category["category_id"]), bands)
        if not isinstance(selected_bands, list) or not selected_bands:
            raise ValueError(f"{category['category_name']} 没有金额区间")
        try:
            distinct_bands = {(Decimal(str(b["min_amount"])), Decimal(str(b["max_amount"]))) for b in selected_bands}
        except (InvalidOperation, TypeError, KeyError) as exc:
            raise ValueError(f"{category['category_name']} 的金额必须是数字") from exc
        if len(distinct_bands) != len(selected_bands):
            raise ValueError(f"{category['category_name']} 的金额区间不能重复")
        if len(selected_bands) > 1048573:
            raise ValueError("查询范围超过 Excel 的行列上限")
    if len({(p["start_date"], p["end_date"]) for p in periods}) != len(periods):
        raise ValueError("日期区间不能重复")
    if len(periods) > 16383:
        raise ValueError("查询范围超过 Excel 的行列上限")
    jobs = []
    for category in categories:
        selected_bands = category_bands.get(str(category["category_id"]), bands)
        for period in periods:
            for band in selected_bands:
                job = {**category, **period, "min_amount": band["min_amount"], "max_amount": band["max_amount"]}
                build_payload(job)
                jobs.append(job)
    return jobs


def result_key(config):
    return (str(config["category_id"]), date_value(config["start_date"]), date_value(config["end_date"]),
            str(Decimal(str(config["min_amount"])).normalize()),
            str(Decimal(str(config["max_amount"])).normalize()))


class ProgressSaveError(ValueError):
    def __init__(self, recovery_path):
        self.recovery_path = recovery_path
        super().__init__("进度文件被占用，最新数据已保存在临时进度中；可稍后继续未完成采集。")


def replace_file(source, destination):
    """Retry Windows sharing/access conflicts without truncating the old file."""
    for delay in (0.05, 0.1, 0.2, 0.4, 0.8, 1, 1, 1, None):
        try:
            source.replace(destination)
            return
        except OSError as exc:
            if not (isinstance(exc, PermissionError) or getattr(exc, 'winerror', None) in (5, 32, 33)) or delay is None:
                raise
            time.sleep(delay)


def save_progress(path, settings, results, status="采集中", error=None):
    payload = {"settings": settings, "results": results, "status": status, "error": error}
    temporary = path.with_suffix(path.suffix + ".tmp")
    with temporary.open('w', encoding='utf-8') as handle:
        handle.write(json.dumps(payload, ensure_ascii=False, indent=2))
        handle.flush()
        os.fsync(handle.fileno())
    try:
        replace_file(temporary, path)
    except OSError as exc:
        raise ProgressSaveError(temporary) from exc


def load_progress(path):
    """Use the newest valid checkpoint, including a failed atomic replacement."""
    candidates = [p for p in (path, path.with_suffix(path.suffix + '.tmp')) if p.is_file()]
    for candidate in sorted(candidates, key=lambda p: p.stat().st_mtime_ns, reverse=True):
        try:
            saved = json.loads(candidate.read_text(encoding='utf-8'))
            check_previous(saved['settings'], saved['results'])
            return saved
        except (ValueError, KeyError, TypeError, OSError, InvalidOperation):
            continue
    raise ValueError("未找到可恢复的有效进度文件")


def check_previous(settings, results):
    expected = {result_key(job) for job in make_plan(settings)}
    seen = set()
    if not isinstance(results, list):
        raise ValueError("进度文件中的结果必须为列表")
    for result in results:
        key = result_key(result["conditions"])
        if key not in expected or key in seen:
            raise ValueError("进度文件包含条件不匹配或重复的结果")
        if type(result["coverage"]) is not int or result["coverage"] < 0:
            raise ValueError("进度文件包含无效人数")
        seen.add(key)


def export_excel(progress_path, output_path, preview=False):
    """Write plain Excel cells without a machine-specific Node/Codex runtime."""
    from openpyxl import Workbook

    saved = load_progress(Path(progress_path))
    settings, results = saved["settings"], saved["results"]
    lookup = {result_key(item["conditions"]): item["coverage"] for item in results}
    workbook = Workbook()
    workbook.remove(workbook.active)
    names = set()
    for category in settings["categories"]:
        bands = settings.get("category_price_ranges", {}).get(str(category["category_id"]), settings["price_ranges"])
        base = re.sub(r"[\\/*?:\[\]]", "_", category["category_name"]).strip("'")[:31] or "类目"
        name = base
        suffix = 2
        while name.lower() in names:
            tail = f"_{suffix}"
            name = base[:31 - len(tail)] + tail
            suffix += 1
        names.add(name.lower())
        sheet = workbook.create_sheet(name)
        sheet.cell(1, 1, "价格区间")
        for col, period in enumerate(settings["periods"], 2):
            sheet.cell(1, col, period["start_date"])
        for row, band in enumerate(bands, 2):
            sheet.cell(row, 1, f"{band['min_amount']}~{band['max_amount']}")
            for col, period in enumerate(settings["periods"], 2):
                key = result_key({**category, **band, **period})
                sheet.cell(row, col, lookup.get(key, "未查询"))
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix=".dmp-export-", suffix=".xlsx", dir=output_path.parent,
                                     delete=False) as handle:
        temporary = Path(handle.name)
    try:
        workbook.save(temporary)
        workbook.close()
        with zipfile.ZipFile(temporary) as archive:
            if archive.testzip() is not None or "xl/workbook.xml" not in archive.namelist():
                raise ValueError("Excel 文件校验失败")
        replace_file(temporary, output_path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    return output_path
