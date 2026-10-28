import json
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

from collector import (BIClient,CollectionError,LoginRequired,Selection,Stopped,
                       checked_response,parse_bindings,period_params,validate_categories,needs_analysis_navigation,group_category_tree)
from excel_export import cell_value,export_result

LIMITS={'minDayDate':'2026-08-27','maxDayDate':'2026-10-25',
        'minWeekDate':'2024-01-13','maxWeekDate':'2026-10-24',
        'minMonthDate':'202401','maxMonthDate':'202609'}


class ContractTests(unittest.TestCase):
    def test_flat_category_tree_is_split_by_actual_parent(self):
        rows=[{'id':'1','pid':'0','name':'一级'},{'id':'2','pid':'1','name':'二级'},
              {'id':'3','pid':'2','name':'三级'},{'id':'4','pid':'1','name':'另一二级'}]
        groups=group_category_tree(rows,'1')
        self.assertEqual([x['id'] for x in groups['1']],['2','4'])
        self.assertEqual([x['id'] for x in groups['2']],['3'])
        self.assertEqual(groups['3'],[])
        for invalid in [[{'id':'2','pid':'9','name':'外部'}],
                        [{'id':'2','pid':'3','name':'循环'},{'id':'3','pid':'2','name':'循环'}]]:
            with self.assertRaises(CollectionError):
                group_category_tree(invalid,'1')

    def test_login_workbench_redirect(self):
        self.assertTrue(needs_analysis_navigation('https://aokesi.ecbis.cn/workTable'))
        self.assertTrue(needs_analysis_navigation('https://aokesi.ecbis.cn/workTable/home'))
        for url in ['https://aokesi.ecbis.cn/login','https://aokesi.ecbis.cn/workTable/productDevelopment',
                    'https://other.example/workTable','http://127.0.0.1:18766/']:
            self.assertFalse(needs_analysis_navigation(url))

    def selection(self,period='7',when=''):
        return Selection('50002768','个人护理/保健/按摩器材','50002768',period,when)

    def test_rolling_parameters_differ_by_endpoint(self):
        p=period_params(self.selection('30'),LIMITS)
        self.assertEqual(p,{'root_date':30,'keyword_date':'20261025','dates_type':30,
                            'start':'2026-09-26','end':'2026-10-25'})
        self.assertEqual(period_params(self.selection('7'),LIMITS)['start'],'2026-10-19')

    def test_week_is_sunday_to_saturday(self):
        p=period_params(self.selection('week','2026-10-21'),LIMITS)
        self.assertEqual((p['start'],p['end'],p['root_date']),('2026-10-18','2026-10-24','20261024'))
        with self.assertRaises(CollectionError):
            period_params(self.selection('week','2026-10-25'),LIMITS)

    def test_month_and_bounds(self):
        self.assertEqual(period_params(self.selection('month','2024-02'),LIMITS)['end'],'2024-02-29')
        for month in ['2026-10','2023-12','2026-13','2026-1','bad']:
            with self.subTest(month=month),self.assertRaises(ValueError):
                period_params(self.selection('month',month),LIMITS)

    def test_failures_are_not_empty_success(self):
        for result in [None,{'status':200,'success':False,'code':0},
                       {'status':500,'success':True,'code':0},
                       {'status':200,'success':True,'code':99}]:
            with self.assertRaises(CollectionError):
                checked_response(result)
        with self.assertRaises(LoginRequired):
            checked_response({'status':200,'success':False,'code':401})
        with self.assertRaises(CollectionError):
            checked_response({'status':429})
        self.assertEqual(checked_response({'status':200,'success':True,'code':0,'data':[]}),[])

    def test_category_validation(self):
        self.assertEqual(validate_categories([{'id':1,'name':'个护'}])[0]['id'],'1')
        for rows in [{},[{'id':'../1','name':'bad'}],[{'id':1,'name':'A'},{'id':1,'name':'B'}]]:
            with self.assertRaises(CollectionError):
                validate_categories(rows)

    def test_current_bindings_and_changed_exports(self):
        folder=Path(__file__).parent/'analysis/assets'
        b=parse_bindings((folder/'demandInsightApi-BMJV-Lhl.js').read_text(encoding='utf-8'),
                         (folder/'shared-Cvf4f1cB.js').read_text(encoding='utf-8'))
        self.assertEqual((b['get'],b['post'],b['crypto']),('br','Tr','hl'))
        demand='import{AA as q,BB as r}from"./shared-new.js";var a=x=>q(`/demand/rootKeyList`,x),b=x=>r(`/demand/getKeyWordByRoot`,x);'
        shared='var ab=store(`wasm`,{});export{ab as Decode};'
        self.assertEqual(parse_bindings(demand,shared),{'get':'AA','post':'BB','crypto':'Decode','shared_filename':'shared-new.js'})
        with self.assertRaises(CollectionError):
            parse_bindings('<html>login</html>',shared)

    def test_cookies_are_not_exported_or_sent_by_python(self):
        class Tab:
            def cookies(self,**kwargs):
                return [{'name':'session','value':'test-secret','domain':'.ecbis.cn'},
                        {'name':'other','value':'private','domain':'.unrelated.example'}]
        c=BIClient()
        c.tab=Tab()
        self.assertIsNone(c.read_cookies())
        self.assertNotIn('test-secret',repr(c.__dict__))

    def test_stop_before_any_request(self):
        stop=threading.Event();stop.set()
        with self.assertRaises(Stopped):
            BIClient(stop=stop).request('/demand/rootKeyList',{})


class ExportTests(unittest.TestCase):
    def result(self):
        return {'status':'complete','selection':{'category_name':'个护'},
                'period':{'start':'2026-10-19','end':'2026-10-25'},
                'source':'https://biplugs.ecbis.cn/prod/demand','fetched_at':'2026-10-27',
                'roots':{'keyWordVoList':[{'rootWord':'吹风机','searchPeople':'1658373','searchPeopleGrow':'0.1145','newRoot':False,'demand':[0]}]},
                'keywords':{'hotWordList':[{'keyWord':'=1+1','rank':'1','searchPeople':'459998','clickPeople':'229778','clickRate':'0.4995','dealPeople':'50142','dealRate':'0.2182','extra':'kept'}]}}

    def test_numeric_precision_and_missing(self):
        self.assertEqual(cell_value('searchPeople','0'),0)
        self.assertIsNone(cell_value('searchPeople',None))
        self.assertEqual(cell_value('searchPeople','--'),'--')
        self.assertEqual(cell_value('searchPeople','123456789012345678'),'123456789012345678')
        self.assertEqual(cell_value('rootWord','123'),'123')

    def test_export_types_extra_fields_and_no_formulas(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder:
            p=export_result(self.result(),Path(folder)/'result.xlsx')
            w=load_workbook(p)
            try:
                s=w['关键词分析']
                self.assertEqual(s['A8'].value,'=1+1')
                self.assertEqual(s['A8'].data_type,'s')
                self.assertEqual(s['C8'].value,459998)
                self.assertEqual(s['E8'].number_format,'0.00%')
                self.assertEqual(s.cell(8,s.max_column).value,'kept')
                self.assertEqual(s.freeze_panes,'B8')
                self.assertEqual(w['词根归类']['G8'].value,False)
                self.assertEqual(len(s.tables),1)
            finally:
                w.close()

    def test_failed_export_preserves_previous_file(self):
        with tempfile.TemporaryDirectory() as folder:
            p=Path(folder)/'result.xlsx';p.write_bytes(b'previous')
            value=self.result();value['keywords']['hotWordList'][0]['extra']='x'*32768
            with self.assertRaises(ValueError):
                export_result(value,p)
            self.assertEqual(p.read_bytes(),b'previous')

    def test_collect_preserves_partial_after_failure(self):
        c=BIClient(log=lambda _:None)
        selection=Selection('50002768','个护','50002768','7')
        root={'keyWordVoList':[],'useAi':True,'minMonth':'2026-10-19','maxMonth':'2026-10-25'}
        calls=[]
        def request(endpoint,params,method='GET'):
            calls.append((endpoint,params))
            if endpoint.endswith('getKeywordAnalysisTime'):
                return LIMITS
            if endpoint.endswith('rootKeyList'):
                return root
            raise LoginRequired('expired')
        c.request=request
        with tempfile.TemporaryDirectory() as folder,patch.object(c.stop,'wait',return_value=False):
            with self.assertRaises(LoginRequired):
                c.collect(selection,folder)
            saved=list(Path(folder).glob('*.json'))
            self.assertEqual(len(saved),1)
            self.assertEqual(json.loads(saved[0].read_text(encoding='utf-8'))['status'],'partial')
            self.assertFalse(list(Path(folder).glob('*.xlsx')))
            self.assertEqual(calls[-1][1]['date'],'20261025')
            self.assertEqual(calls[-1][1]['useAi'],True)


if __name__=='__main__':
    unittest.main()
