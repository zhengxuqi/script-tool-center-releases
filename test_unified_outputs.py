import io
import json
import os
import sys
import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / 'app/versions/1.0.6'))
from test_hub import seed, eventually
from hubcore.runtime import TaskQueue
from hubcore.dmp_portal import DMPPortal
from tool_sdk import atomic_json


class UnifiedOutputTests(unittest.TestCase):
    def test_exports_share_one_folder_and_same_names_do_not_overwrite(self):
        with tempfile.TemporaryDirectory() as temp:
            store = seed(Path(temp))
            old = store.outputs / '相同名称.xlsx'
            old.write_bytes(b'keep-old')
            names = [store.save_export(tool, '相同名称.xlsx', data) for tool, data in (
                ('boundary', b'boundary'), ('dmp', b'dmp'), ('dmp', b'new-dmp'))]
            self.assertEqual(len(set(names)), 3)
            self.assertEqual(old.read_bytes(), b'keep-old')
            for name, value in zip(names, (b'boundary', b'dmp', b'new-dmp')):
                self.assertEqual((store.outputs / name).read_bytes(), value)
                self.assertEqual(Path(name).name, name)
            rows = {r['path']: r['tool_id'] for r in store.results()}
            self.assertEqual([rows[name] for name in names], ['boundary', 'dmp', 'dmp'])
            self.assertFalse(list(store.outputs.glob('*.tmp')))

    def test_queue_archives_to_root_and_resume_keeps_internal_checkpoint(self):
        with tempfile.TemporaryDirectory() as temp:
            store = seed(Path(temp))
            env = Mock()
            env.ensure.return_value = Path(sys.executable)
            queue = TaskQueue(store, env)
            job = queue.submit('dmp', 'collect', {'steps': 1})
            eventually(lambda: store.job(job['id'])['status'] == 'completed')
            result = store.job(job['id'])['result']
            self.assertEqual(len(result['files']), 1)
            self.assertTrue((store.outputs / result['files'][0]).is_file())
            workspace = store.state / 'jobs' / job['id'] / 'output'
            self.assertTrue((workspace / 'result.json').is_file())
            self.assertFalse((store.outputs / 'dmp').exists())
            checkpoint = workspace / 'progress.json'
            atomic_json(checkpoint, {'saved': True})
            resumed = queue.submit('dmp', 'collect', {}, resume=job['id'])
            self.assertEqual(Path(resumed['resume_path']), checkpoint)
            eventually(lambda: not store.unfinished())
            self.assertTrue(checkpoint.is_file())
            legacy_id = 'old-job'
            old = {**store.job(job['id']), 'id': legacy_id}
            store.put_job(old)
            legacy = store.outputs / 'dmp' / legacy_id / 'progress.json'
            atomic_json(legacy, {'legacy': True})
            resumed = queue.submit('dmp', 'collect', {}, resume=legacy_id)
            self.assertEqual(Path(resumed['resume_path']), legacy)
            eventually(lambda: not store.unfinished())

    def test_dmp_finishes_into_shared_folder_without_changing_its_excel(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as temp:
            store = seed(Path(temp))
            workspace = store.state / 'tools/dmp/workspace'
            workspace.mkdir(parents=True)
            book = Workbook()
            book.active.title = '电吹风'
            book.active.append(['金额区间', '2026-09'])
            book.active.append(['0～100', 123])
            source = workspace / '我的覆盖人数.xlsx'
            book.save(source)
            atomic_json(workspace / '覆盖人数.progress.json', {'saved': True})
            job = {'id': 'managed-job', 'tool_id': 'dmp', 'action': 'collect', 'created': 1,
                   'status': 'running', 'params': {'_portal_ticket': 'a' * 32}, 'logs': []}
            store.put_job(job)
            portal = DMPPortal.__new__(DMPPortal)
            portal.store = store
            portal.tasks = SimpleNamespace(gate=threading.RLock())
            portal.json = Mock(side_effect=[{}, {'phase': 'finished', 'task': {
                'status': '完成', 'completed': 1, 'total': 1, 'output': source.name, 'error': None}}])
            portal.run_job(job)
            saved = store.job(job['id'])
            self.assertEqual(saved['status'], 'completed')
            exported = store.outputs / saved['result']['filename']
            self.assertEqual(exported.parent, store.outputs)
            self.assertEqual(exported.read_bytes(), source.read_bytes())
            self.assertEqual(list(load_workbook(exported).active.values), [('金额区间', '2026-09'), ('0～100', 123)])
            self.assertEqual(len(saved['result']['files']), 2)
            self.assertTrue((workspace / '覆盖人数.progress.json').is_file())


if __name__ == '__main__':
    unittest.main()
